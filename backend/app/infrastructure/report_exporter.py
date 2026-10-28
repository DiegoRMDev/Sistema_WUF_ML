"""Generación de archivos Excel y PDF para HU05."""

from __future__ import annotations

from collections import Counter
from datetime import datetime
from html import escape
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table as PdfTable, TableStyle


HEADERS = [
    "SKU", "Nombre", "Especie", "Sexo", "Edad", "Tamaño", "Condición especial",
    "Orden de publicación", "Probabilidad estimada de adopción", "Prioridad",
    "Atributos desfavorecidos",
]


def _age(months: int | None, estimated: bool | None) -> str:
    if months is None:
        return "No indicada"
    years, remaining = divmod(int(months), 12)
    if years:
        value = f"{years} {'año' if years == 1 else 'años'}"
        if remaining:
            value += f" {remaining} {'mes' if remaining == 1 else 'meses'}"
    else:
        value = f"{remaining} {'mes' if remaining == 1 else 'meses'}"
    return f"{value} (est.)" if estimated else value


def _condition(item: dict[str, Any]) -> str:
    if not item["condicion_especial"]:
        return "Ninguna"
    return item.get("condicion_detalle") or "Condición especial"


def _disadvantaged_attributes(item: dict[str, Any]) -> str:
    """Resume los criterios detectados y conserva visible la edad desconocida."""
    criteria = [
        f"{criterion['nombre']}{' (est.)' if criterion['estimado'] else ''}"
        for criterion in item.get("criterios_desfavorecidos", [])
    ]
    if item.get("edad_desconocida"):
        criteria.append("Edad no indicada")
    return ", ".join(criteria) if criteria else "Ninguno"


def _excel_text(value: Any) -> str:
    """Fuerza texto en Excel y bloquea fórmulas que pudieran venir de los CSV."""
    text = str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _rows(items: list[dict[str, Any]], *, excel_safe: bool = False) -> list[list[Any]]:
    as_text = _excel_text if excel_safe else str
    return [[
        as_text(item["sku"]), as_text(item["nombre"]), as_text(str(item["especie"]).capitalize()),
        as_text("Hembra" if str(item["sexo"]).upper() == "H" else "Macho"),
        as_text(_age(item.get("edad_meses"), item.get("edad_estimada"))),
        as_text(str(item["tamano"]).capitalize()), as_text(_condition(item)), int(item["orden_publicacion"]),
        float(item["probabilidad_adopcion"]), str(item["nivel_prioridad"]).capitalize(),
        as_text(_disadvantaged_attributes(item)),
    ] for item in items]


def _priority_summary(items: list[dict[str, Any]]) -> str:
    counts = Counter(str(item["nivel_prioridad"]).casefold() for item in items)
    return f"Prioridad alta: {counts['alta']} · media: {counts['media']} · baja: {counts['baja']} · total: {len(items)}"


def export_xlsx(report: dict[str, Any], generated_at: datetime) -> bytes:
    """Crea un libro con el periodo, el modelo aplicado y las filas ordenadas del reporte."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Priorización"
    sheet.append(["Reporte semanal de priorización"])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="28344F")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    capture = report["captura"]
    model = report["modelo"]
    sheet.append([f"Semana ISO {capture['semana']} · Captura {capture['fecha_captura'].strftime('%d/%m/%Y')} · Modelo v{model['version']} · Generado {generated_at.strftime('%d/%m/%Y %H:%M')} (Lima)"])
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
    sheet["A2"].font = Font(italic=True, color="6B6F76")
    sheet.append([_priority_summary(report["items"])])
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(HEADERS))
    sheet["A3"].font = Font(bold=True, color="28344F")
    sheet.append(["Segmento desfavorecido: edad avanzada (96 meses o más), tamaño grande o condición especial. Los valores estimados se marcan; una edad no indicada no cumple el criterio etario."])
    sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=len(HEADERS))
    sheet["A4"].font = Font(italic=True, color="6B6F76")
    sheet.append(HEADERS)
    for cell in sheet[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="465472")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[5].height = 32
    for values in _rows(report["items"], excel_safe=True):
        sheet.append(values)
        sheet.cell(sheet.max_row, 9).number_format = "0.0%"
    widths = [14, 24, 14, 12, 20, 14, 34, 19, 28, 14, 38]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A6"
    if report["items"]:
        last_column = get_column_letter(len(HEADERS))
        table = Table(displayName="PrioritizacionSemanal", ref=f"A5:{last_column}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_pdf(report: dict[str, Any], generated_at: datetime) -> bytes:
    """Construye un PDF apaisado para impresión y lectura de la tabla completa."""
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=landscape(letter), rightMargin=0.35 * inch,
        leftMargin=0.35 * inch, topMargin=0.35 * inch, bottomMargin=0.35 * inch,
        title="Reporte semanal de priorización",
    )
    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    title_style.textColor = colors.HexColor("#28344F")
    meta_style = ParagraphStyle("ReportMeta", parent=styles["Normal"], fontSize=8.5, leading=12)
    cell_style = ParagraphStyle("ReportCell", parent=styles["Normal"], fontSize=7, leading=9, alignment=TA_LEFT)
    right_style = ParagraphStyle("ReportNumber", parent=cell_style, alignment=TA_RIGHT)
    capture = report["captura"]
    model = report["modelo"]
    story = [
        Paragraph("Reporte semanal de priorización", title_style),
        Paragraph(
            f"Semana ISO {capture['semana']} · Captura {capture['fecha_captura'].strftime('%d/%m/%Y')}"
            f" · Modelo v{escape(str(model['version']))} · Generado {generated_at.strftime('%d/%m/%Y %H:%M')} (Lima)",
            meta_style,
        ),
        Paragraph(escape(_priority_summary(report["items"])), meta_style),
        Paragraph(
            "Segmento desfavorecido: edad avanzada (96 meses o más), tamaño grande o condición especial. "
            "Los valores estimados se marcan; una edad no indicada no cumple el criterio etario.",
            meta_style,
        ),
        Spacer(1, 10),
    ]
    data: list[list[Any]] = [[Paragraph(escape(header), cell_style) for header in HEADERS]]
    for item in _rows(report["items"]):
        formatted = list(item)
        formatted[8] = f"{float(item[8]) * 100:.1f} %"
        data.append([
            Paragraph(escape(str(value)), right_style if index in {7, 8} else cell_style)
            for index, value in enumerate(formatted)
        ])
    available_width = landscape(letter)[0] - 0.7 * inch
    fractions = [0.055, 0.085, 0.055, 0.045, 0.075, 0.055, 0.12, 0.07, 0.12, 0.055, 0.265]
    column_widths = [available_width * fraction / sum(fractions) for fraction in fractions]
    table = PdfTable(data, colWidths=column_widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#28344F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D8DCE4")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F5F7")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    document.build(story)
    return output.getvalue()


